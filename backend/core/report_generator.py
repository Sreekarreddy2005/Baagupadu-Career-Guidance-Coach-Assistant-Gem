import io
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_persona_docx(profile: dict) -> io.BytesIO:
    doc = Document()
    
    # Title
    title = doc.add_heading('Baagupadu: Your Authentic Persona & Roadmap', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph('\n')
    
    # 1. Demographics
    demographics = profile.get('life_stage_data', {}).get('demographics', {})
    if demographics:
        doc.add_heading('1. Demographics', level=1)
        p = doc.add_paragraph()
        p.add_run('Name: ').bold = True
        p.add_run(f"{demographics.get('full_name', 'N/A')}\n")
        p.add_run('Desired Role: ').bold = True
        p.add_run(f"{demographics.get('desired_role', 'N/A')}\n")
        p.add_run('Primary Goal: ').bold = True
        p.add_run(f"{demographics.get('primary_goal', 'N/A')}\n")
        doc.add_paragraph('\n')
        
    # 2. Persona Synthesis
    persona = profile.get('persona', {})
    if persona:
        doc.add_heading('2. Your Authentic Self', level=1)
        
        core_identity = persona.get('core_identity', {})
        if core_identity:
            p = doc.add_paragraph()
            p.add_run('Core Identity: ').bold = True
            p.add_run(f"{core_identity.get('archetype_name', 'Unknown Archetype')}\n")
            p.add_run('Tagline: ').bold = True
            p.add_run(f"{core_identity.get('tagline', '')}\n\n")
            doc.add_paragraph(core_identity.get('description_for_user', core_identity.get('description', '')))
            
        strengths = persona.get('strengths', [])
        if strengths:
            doc.add_heading('Key Strengths', level=2)
            for strength in strengths:
                p = doc.add_paragraph(style='List Bullet')
                p.add_run(f"{strength.get('trait', '')}: ").bold = True
                p.add_run(strength.get('evidence', ''))
                
        growth = persona.get('growth_areas', [])
        if growth:
            doc.add_heading('Growth Areas', level=2)
            for area in growth:
                p = doc.add_paragraph(style='List Bullet')
                p.add_run(f"{area.get('area', '')}: ").bold = True
                p.add_run(area.get('compassionate_framing', ''))

        shadow = persona.get('shadow_traits', [])
        if shadow:
            doc.add_heading('Shadow Traits', level=2)
            for trait in shadow:
                p = doc.add_paragraph(style='List Bullet')
                p.add_run(f"{trait.get('trait', '')}: ").bold = True
                p.add_run(trait.get('acknowledgment', ''))
                
        doc.add_paragraph('\n')
        
    # 3. Career Roadmap (Guidance)
    roadmap = profile.get('guidance', {}).get('roadmap', {})
    if roadmap:
        doc.add_heading('3. Career Roadmap', level=1)
        
        primary_path = roadmap.get('primary_career_path', {})
        if primary_path:
            p = doc.add_paragraph()
            p.add_run('Recommended Career Path: ').bold = True
            p.add_run(f"{primary_path.get('title', 'Unknown')}\n")
            doc.add_paragraph(primary_path.get('reasoning', ''))
            
        milestones = roadmap.get('action_plan', [])
        if milestones:
            doc.add_heading('Action Plan', level=2)
            for period in milestones:
                if isinstance(period, dict):
                    timeframe = period.get('timeframe', 'Action')
                    tasks = period.get('tasks', [])
                    
                    doc.add_heading(timeframe, level=3)
                    for task in tasks:
                        action = task.get('action', '')
                        points = task.get('points', 0)
                        p = doc.add_paragraph(style='List Bullet')
                        p.add_run(action)
                        if points:
                            p.add_run(f" (+{points} SP)").italic = True
                else:
                    doc.add_paragraph(str(period), style='List Bullet')

    doc.add_paragraph('\n')

    # 4. Final Summary & Quality Signals
    final_summary = roadmap.get('final_summary', {})
    if final_summary:
        doc.add_heading('4. Coach Summary', level=1)
        doc.add_paragraph(final_summary.get('coaching_feedback', ''))
        
        insights = final_summary.get('deep_insights', [])
        if insights:
            doc.add_heading('Deep Insights', level=2)
            for insight in insights:
                doc.add_paragraph(insight, style='List Bullet')

    quality = roadmap.get('quality_signals', {})
    if quality:
        doc.add_heading('Session Quality Metrics', level=2)
        p = doc.add_paragraph()
        for k, v in quality.items():
            p.add_run(f"{k.replace('_', ' ').title()}: ").bold = True
            p.add_run(f"{v}\n")

    doc.add_paragraph('\n')
    footer = doc.add_paragraph()
    footer_run = footer.add_run('Generated by Baagupadu AI Coach (Sahayam)')
    footer_run.italic = True
    footer_run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    
    return file_stream
